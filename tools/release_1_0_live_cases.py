"""R10 observed live scenarios; invoked only by the isolated native worker.

These handlers perform real signed HTTP requests. Their model-free tests verify
protocol/oracle behavior, and cannot certify the isolated-live layer.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
import secrets
import time
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from typing import Any

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_ARTIFACT_BYTES = 4 << 20
MAX_JSON_BYTES = 8 << 20
_RAW_ID = re.compile(r"raw_[0-9a-f]{16}")
_MESSAGE_ID = re.compile(r"msg_[0-9a-f]{16}")


class CaseProtocolError(RuntimeError):
    """Closed protocol/observation failure, with no raw response in its text."""


@dataclass(frozen=True)
class WordFixture:
    case_id: str
    filename: str
    output_filename: str
    source_ref: str
    message_ref: str
    telegram_message_id: int
    content: bytes
    prompt: str
    required_text: tuple[str, ...]

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.content).hexdigest()


WORD_VARIANTS = (
    ("R10-LIVE-DOC-WORD-FIRST-GEN", "brief.txt", "brief.docx"),
    ("R10-LIVE-DOC-WORD-NAMED-CYRILLIC", "материалы кухни.txt", "итог кухни.docx"),
    ("R10-LIVE-DOC-WORD-NAMED-VERSION", "заметки-v2.txt", "отчёт-v2.docx"),
)


def word_fixture(case_id: str, run_id: str) -> WordFixture:
    if re.fullmatch(r"[0-9a-f]{32}", run_id) is None:
        raise CaseProtocolError("run_identity_invalid")
    variant = next((row for row in WORD_VARIANTS if row[0] == case_id), None)
    if variant is None:
        raise CaseProtocolError("word_case_unknown")
    index = WORD_VARIANTS.index(variant)
    _, filename, output_name = variant
    token = hashlib.sha256(f"{run_id}:{case_id}".encode()).hexdigest()[:16]
    facts = (
        f"Проект: КУХНЯ-{token}",
        f"Ответственный: Соколова-{index + 1}",
        f"Количество участников: {17 + index}",
    )
    prompts = (
        f"Создай Word-файл {output_name} по приложенному {filename}. "
        "Перенеси все три строки исходника без изменения значений.",
        f"Возьми данные из файла «{filename}» и подготовь документ «{output_name}» в формате DOCX. "
        "Сохрани в документе каждую из трёх строк и все её значения.",
        f"Из приложенного {filename} сделай {output_name}. Нужен скачиваемый Word-документ, "
        "содержащий все три исходные строки с точными значениями.",
    )
    return WordFixture(
        case_id,
        filename,
        output_name,
        f"telegram-file:r10-{token}",
        f"r10-message:{token}",
        int(token[:12], 16),
        ("\n".join(facts) + "\n").encode(),
        prompts[index],
        facts,
    )


class SignedSession:
    """One actor and one first-turn submission; raw evidence stays private."""

    def __init__(self, client: Any, *, bridge_secret: str, chat_id: int) -> None:
        if not bridge_secret or type(chat_id) is not int or chat_id <= 0:
            raise CaseProtocolError("signed_case_actor_invalid")
        self.client = client
        self.secret = bridge_secret
        self.chat_id = str(chat_id)
        self.evidence: list[dict[str, Any]] = []
        self.artifacts: dict[str, bytes] = {}
        self.chat_submissions = 0

    def _request(self, method: str, path: str, payload: Mapping[str, Any] | None = None) -> Any:
        from friday.security import sign_bridge_request

        body = (
            json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
            if payload is not None
            else b""
        )
        timestamp = int(time.time())
        nonce = secrets.token_hex(16)
        signature = sign_bridge_request(
            self.secret,
            timestamp=timestamp,
            method=method,
            path=path,
            external_user_id=self.chat_id,
            chat_id=self.chat_id,
            nonce=nonce,
            body=body,
        )
        headers = {
            "Content-Type": "application/json",
            "X-Friday-Timestamp": str(timestamp),
            "X-Friday-User": self.chat_id,
            "X-Friday-Chat": self.chat_id,
            "X-Friday-Nonce": nonce,
            "X-Friday-Signature": signature,
        }
        response = self.client.request(
            method, path, content=body or None, headers=headers, follow_redirects=False
        )
        self.evidence.append(
            {"method": method, "path": path, "request": payload, "status": response.status_code}
        )
        return response

    def json(self, method: str, path: str, payload: Mapping[str, Any] | None = None) -> Mapping[str, Any]:
        response = self._request(method, path, payload)
        if response.status_code != 200:
            raise CaseProtocolError("case_http_status")
        if len(response.content) > MAX_JSON_BYTES:
            raise CaseProtocolError("case_json_oversized")
        try:
            value = response.json()
        except (ValueError, UnicodeError):
            raise CaseProtocolError("case_json_invalid") from None
        if not isinstance(value, dict):
            raise CaseProtocolError("case_json_invalid")
        self.evidence[-1]["response"] = value
        return value

    def first_chat(self, fixture: WordFixture) -> Mapping[str, Any]:
        if self.chat_submissions:
            raise CaseProtocolError("case_first_attempt_already_submitted")
        # Count before dispatch, including a failed HTTP attempt. There is no retry.
        self.chat_submissions += 1
        from tools.document_contour_live_battery import Harness

        return self.json(
            "POST",
            "/api/chat",
            {
                "message": fixture.prompt,
                "source_ref": fixture.message_ref,
                "telegram_message_id": fixture.telegram_message_id,
                "telegram_user": {"id": int(self.chat_id), "first_name": "R10", "language_code": "ru"},
                "enable_tools": True,
                "document": Harness.document(
                    fixture.filename, "text/plain", fixture.content, fixture.source_ref
                ),
            },
        )

    def download(self, raw_id: str) -> bytes:
        if not isinstance(raw_id, str) or _RAW_ID.fullmatch(raw_id) is None:
            raise CaseProtocolError("artifact_handle_invalid")
        response = self._request("GET", f"/api/files/{raw_id}")
        if response.status_code != 200:
            raise CaseProtocolError("artifact_download_failed")
        payload = response.content
        if not payload or len(payload) > MAX_ARTIFACT_BYTES:
            raise CaseProtocolError("artifact_bytes_invalid")
        digest = hashlib.sha256(payload).hexdigest()
        self.artifacts[digest] = payload
        self.evidence[-1].update(sha256=digest, size_bytes=len(payload))
        return payload


def _docx_text(payload: bytes) -> str:
    # Independently parse delivered bytes. Bound compressed and expanded input
    # before the reader opens OOXML; body strings in the chat are not an oracle.
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos = archive.infolist()
            if (
                not infos
                or len(infos) > 1024
                or len({item.filename for item in infos}) != len(infos)
                or sum(item.file_size for item in infos) > 16 << 20
                or any(item.flag_bits & 1 for item in infos)
            ):
                raise CaseProtocolError("docx_package_invalid")
            if "word/document.xml" not in archive.namelist():
                raise CaseProtocolError("docx_package_invalid")
        from docx import Document

        document = Document(io.BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return "\n".join(parts)
    except CaseProtocolError:
        raise
    except Exception as exc:
        # A damaged package is an observed oracle failure, never success.
        raise CaseProtocolError("docx_reader_failed") from exc


def _metadata(row: Mapping[str, Any]) -> Mapping[str, Any]:
    try:
        value = json.loads(str(row.get("metadata_json") or "{}"))
    except (ValueError, TypeError):
        return {}
    return value if isinstance(value, dict) else {}


def _require_owner_actor(identity: Mapping[str, Any], session: SignedSession, storage: Any) -> str:
    """Bind the observed owner role to the exact signed Telegram principal."""

    from friday.permissions import LEGACY_OWNER_USER_ID

    actor = identity.get("actor")
    user = identity.get("user")
    if (
        not isinstance(actor, Mapping)
        or actor.get("user_id") != LEGACY_OWNER_USER_ID
        or actor.get("preset_key") != "owner"
        or actor.get("source") != "telegram-bridge"
        or not isinstance(user, Mapping)
        or user.get("id") != LEGACY_OWNER_USER_ID
        or user.get("preset_key") != "owner"
        or user.get("status") != "active"
    ):
        raise CaseProtocolError("case_owner_actor_mismatch")
    try:
        principal_owner = storage.resolve_identity("telegram", session.chat_id)
    except (TypeError, ValueError):
        principal_owner = None
    if principal_owner != LEGACY_OWNER_USER_ID:
        raise CaseProtocolError("case_owner_principal_mismatch")
    return LEGACY_OWNER_USER_ID


def run_word_case(session: SignedSession, storage: Any, fixture: WordFixture) -> dict[str, Any]:
    started = time.monotonic()
    failures: list[str] = []
    observed: dict[str, Any] = {"fixture_sha256": fixture.sha256}
    try:
        identity = session.json("GET", "/api/me")
        user_id = _require_owner_actor(identity, session, storage)
        reply = session.first_chat(fixture)
        message_id = reply.get("message_id")
        conversation_id = reply.get("conversation_id")
        if (
            not isinstance(message_id, str)
            or _MESSAGE_ID.fullmatch(message_id) is None
            or not isinstance(conversation_id, str)
            or not conversation_id
        ):
            raise CaseProtocolError("case_message_identity_missing")
        rows = storage.get_conversation_messages(conversation_id, user_id=user_id, limit=100)
        if (
            sum(row.get("role") == "user" for row in rows) != 1
            or sum(row.get("role") == "assistant" for row in rows) != 1
        ):
            failures.append("case_not_first_conversation_turn")
        source_id = storage.resolve_owned_file_source_ref(user_id, user_id, fixture.source_ref)
        if hashlib.sha256(session.download(source_id)).hexdigest() != fixture.sha256:
            failures.append("source_bytes_mismatch")
        files = reply.get("files")
        if not isinstance(files, list) or len(files) != 1 or not isinstance(files[0], Mapping):
            raise CaseProtocolError("one_required_artifact_missing")
        item = files[0]
        raw_id = item.get("id")
        if raw_id == source_id:
            failures.append("output_is_source_handle")
        if item.get("filename") != fixture.output_filename or item.get("mime_type") != DOCX_MIME:
            failures.append("output_format_or_filename_mismatch")
        if item.get("download_url") != f"/api/files/{raw_id}":
            failures.append("artifact_download_binding_invalid")
        payload = session.download(raw_id)
        digest = hashlib.sha256(payload).hexdigest()
        if (
            item.get("sha256") != digest
            or type(item.get("size_bytes")) is not int
            or item["size_bytes"] != len(payload)
        ):
            failures.append("artifact_descriptor_mismatch")
        try:
            inline = base64.b64decode(item.get("content_base64", ""), validate=True)
        except (ValueError, TypeError):
            inline = b""
        if inline != payload:
            failures.append("inline_download_disagree")
        message = storage.get_message(message_id, user_id) or {}
        descriptors = _metadata(message).get("generated_files")
        wanted = {
            key: item.get(key)
            for key in ("id", "filename", "mime_type", "size_bytes", "sha256", "download_url")
        }
        if (
            message.get("role") != "assistant"
            or message.get("conversation_id") != conversation_id
            or descriptors != [wanted]
        ):
            failures.append("artifact_history_binding_invalid")
        text = _docx_text(payload)
        if not all(fact in text for fact in fixture.required_text):
            failures.append("docx_source_facts_missing")
        observed.update(artifact_sha256=digest, artifact_size_bytes=len(payload))
    except CaseProtocolError as exc:
        failures.append(str(exc))
    return {
        "id": fixture.case_id,
        "status": "FAIL" if failures else "PASS",
        "failure_codes": sorted(set(failures)),
        "attempt": 1,
        "chat_submissions": session.chat_submissions,
        "duration_ms": round((time.monotonic() - started) * 1000),
        "observed_safe": observed,
    }
